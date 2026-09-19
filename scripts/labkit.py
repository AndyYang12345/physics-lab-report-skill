#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
物理实验报告 · 工具库
所有 docx 操作集中在这里，build.py / fill.py 只负责调用。

设计的核心约定：
  · 模板里所有定位都靠模板自带的固定标签（"实验日期："、"实验目的："、"实验名称"），
    绝不靠正文内容匹配 —— 正文是每次生成的，不能当锚点。
  · 段落格式一律 宋体 / 小四(sz=24) / 1.5 倍行距(line=360,auto)；
    图表题注用 五号(sz=21)。封面字段用 三号(sz=32)。
"""
import os
import re
import shutil
import subprocess
from pathlib import Path

# python-docx 允许缺失：这样 `lab.py doctor` 在依赖没装好时也能跑起来告诉你缺什么
try:
    import docx
    from docx.shared import Pt, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError as _e:            # pragma: no cover
    _DOCX_ERR = _e
else:
    _DOCX_ERR = None


def _need_docx():
    if _DOCX_ERR is not None:
        raise RuntimeError('缺少 python-docx，无法读写 docx：pip install python-docx') from _DOCX_ERR

# --------------------------------------------------------------------------
# 路径（均可用环境变量覆盖，便于装到任意位置）
#   PHYSLAB_HOME      工作根目录        默认 ~/物理实验报告
#   PHYSLAB_TEMPLATE  空白模板文件路径   默认 $PHYSLAB_HOME/templates/实验报告模板-空白.docx
# --------------------------------------------------------------------------
def _env_path(name, default):
    v = os.environ.get(name)
    return Path(v).expanduser() if v else default


ROOT     = _env_path('PHYSLAB_HOME', Path.home() / '物理实验报告')
TEMPLATE = _env_path('PHYSLAB_TEMPLATE', ROOT / 'templates' / '实验报告模板-空白.docx')
PROFILE  = ROOT / '.profile'
BUILD    = ROOT / '.build'
SUBMIT   = ROOT / '提交'

# --------------------------------------------------------------------------
# 格式常量
# --------------------------------------------------------------------------
FONT     = '宋体'
SZ_COVER = 32   # 三号  —— 封面字段
SZ_BODY  = 24   # 小四  —— 正文
SZ_CAP   = 21   # 五号  —— 图表题注
SZ_TABLE = 24   # 小四  —— 数据表
LINE     = 1.5  # 正文行距

REQUIRED = ['姓名', '学号', '座位号', '专业', '专业班级',
            '课程名称', '开课学期', '指导老师', '实验时段']


# ==========================================================================
# 个人信息
# ==========================================================================
def read_profile(path=None):
    """读取 .profile → dict。格式 `键: 值`，# 之后为注释。"""
    p = Path(path or PROFILE)
    data = {}
    if not p.exists():
        return data
    for line in p.read_text(encoding='utf-8').splitlines():
        line = line.split('#')[0].strip()
        if not line or ':' not in line:
            continue
        k, v = line.split(':', 1)
        data[k.strip()] = v.strip()
    return data


def missing_fields(prof):
    """返回 .profile 里缺失或仍带 ⚠ 的字段（需要问用户）"""
    bad = [k for k in REQUIRED if not prof.get(k)]
    bad += [k for k in REQUIRED if prof.get(k, '').startswith('⚠')]
    return bad


def safe_name(s):
    """文件名安全化：去掉 / \\ : * ? " < > | 并压掉首尾空格"""
    return re.sub(r'[\\/:*?"<>|\r\n\t]', '', str(s)).strip()


def submission_name(prof, experiment, ext='pdf'):
    """提交文件名：座位号-姓名-实验名称.pdf"""
    return '%s-%s-%s.%s' % (safe_name(prof['座位号']),
                            safe_name(prof['姓名']),
                            safe_name(experiment), ext)


# ==========================================================================
# docx 基础操作
# ==========================================================================
def style_run(run, text=None, size=SZ_BODY, bold=False,
              underline=None, sub=False):
    """统一设置 run 的字体/字号/粗细/下划线/下标，含中文字体(eastAsia)"""
    if text is not None:
        run.text = text
    run.font.name = FONT
    run.font.size = Pt(size / 2)          # docx 用半磅
    run.font.bold = bold
    if underline is not None:
        run.font.underline = underline
    if sub:
        run.font.subscript = True
    # 中文字体必须单独设 eastAsia，否则 Word 里会回落成默认字体
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), FONT)
    return run


def style_para(p, align=None, line=LINE, indent=False, space_after=0):
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    if indent:
        pf.first_line_indent = Pt(SZ_BODY / 2 * 2)   # 首行缩进 2 字符
    return p


def clear_cell(cell):
    """清空单元格里所有段落（保留 tcPr）"""
    for p in list(cell.paragraphs):
        p._element.getparent().remove(p._element)


def cell_para(cell, text, *, bold=False, size=SZ_BODY, align=None,
              indent=False, line=LINE, sub_parts=None):
    """
    往单元格里加一个段落。
    sub_parts: [(文本, 是否下标), ...] —— 需要真下标时用它，别用 Unicode 下标字符
    """
    p = cell.add_paragraph()
    style_para(p, align=align, line=line, indent=indent)
    if sub_parts:
        for t, is_sub in sub_parts:
            style_run(p.add_run(), t, size=size, bold=bold, sub=is_sub)
    else:
        style_run(p.add_run(), text, size=size, bold=bold)
    return p


def underlined_runs(p):
    """按出现顺序返回段落里所有带下划线的 run（信息行的填空位）"""
    return [r for r in p.runs if r.font.underline]


def set_blank(runs, idxs, value):
    """把一组下划线 run 当作一个填空位：首个 run 写值，其余清空"""
    if not idxs or idxs[0] >= len(runs):
        return
    runs[idxs[0]].text = value
    for i in idxs[1:]:
        if i < len(runs):
            runs[i].text = ''


def set_table_borders(table, sz=4):
    """给表格加全框线（python-docx 没有边框 API，只能写 XML）"""
    tblPr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), '000000')
        borders.append(e)
    tblPr.append(borders)


# ==========================================================================
# 模板定位器（全部基于模板自带的固定标签）
# ==========================================================================
def find_para(doc, prefix, limit=0):
    """按文本前缀找段落"""
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            return p
    raise LookupError('模板里找不到以 %r 开头的段落' % prefix)


def find_table_by_cell(doc, text, row=0, col=0):
    """找 cell(row,col) 以 text 开头的表格"""
    for t in doc.tables:
        try:
            if t.cell(row, col).text.strip().startswith(text):
                return t
        except IndexError:
            continue
    raise LookupError('模板里找不到单元格 %r 的表格' % text)


# ==========================================================================
# 组装报告（预习模式）
# ==========================================================================
def _missing_template():
    raise FileNotFoundError(
        '找不到空白模板：%s\n'
        '本仓库不附带任何学校的模板文件（版权原因）。请按 reference/format.md 的说明\n'
        '准备一份空白模板，放到该路径，或设环境变量 PHYSLAB_TEMPLATE 指向它。' % TEMPLATE)


def build_report(spec, out_path, profile=None):
    _need_docx()
    """
    spec = {
      "experiment": "分光计的调整与使用",
      "number":     "",                       # 有就写，没有留空
      "date":       "",                       # 预习阶段留空，实验后再填
      "sections": [
         {"h": "实验目的："},
         {"p": "1．了解分光计的结构……"},
         {"f": "fig1.png", "cap": "图1  分光计结构示意图"},
         ...
      ]
    }
    返回 out_path
    """
    prof = profile or read_profile()
    miss = missing_fields(prof)
    if miss:
        raise ValueError('.profile 缺少字段：%s' % '、'.join(miss))

    doc = docx.Document(str(TEMPLATE)) if TEMPLATE.exists() else _missing_template()

    # ---- 1. 封面 5 个字段 ----
    cover = find_table_by_cell(doc, '课程名称')
    for i, key in enumerate(['课程名称', '姓名', '学号', '专业班级', '开课学期']):
        cell = cover.cell(i, 1)
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_run(p.add_run(), prof[key], size=SZ_COVER, bold=True)

    # ---- 2. 标题：模板里只有“实验 ”，追加名称 ----
    head = None
    for p in doc.paragraphs:
        if p.text.strip() == '实验':
            head = p
            break
    if head is None:
        raise LookupError('模板标题段（文本恰为“实验”）未找到')
    name = spec['experiment']
    if spec.get('number'):
        name = '%s %s' % (spec['number'], name)
    style_run(head.add_run(), name, size=SZ_COVER, bold=True)

    # ---- 3. 信息行 ----
    info1 = find_para(doc, '实验日期：')          # 日期 / 姓名 / 座位号
    set_blank(underlined_runs(info1), [0], spec.get('date') or '')
    set_blank(underlined_runs(info1), [1], ' %s ' % prof['姓名'])
    set_blank(underlined_runs(info1), [2], ' %s ' % prof['座位号'])

    info2 = find_para(doc, '实验时段：')          # 时段 / 专业 / 指导老师
    r2 = underlined_runs(info2)                   # 9 个 run，每 3 个一组
    set_blank(r2, [0, 1, 2], ' %s ' % prof['实验时段'])
    set_blank(r2, [3, 4, 5], ' %s ' % prof['专业'])
    set_blank(r2, [6, 7, 8], ' %s ' % prof['指导老师'])

    # ---- 4. 正文单元格 ----
    body_tbl = find_table_by_cell(doc, '实验目的')
    cell = body_tbl.cell(0, 0)
    clear_cell(cell)
    for item in spec['sections']:
        if 'h' in item:                                   # 小节标题（加粗）
            cell_para(cell, item['h'], bold=True)
        elif 'p' in item:                                 # 正文段
            cell_para(cell, item['p'])
        elif 'f' in item:                                 # 插图 + 题注
            _insert_figure(cell, item['f'], item.get('cap', ''),
                           item.get('width_cm', 12.6))
        elif 'blank' in item:                             # 空行
            cell.add_paragraph()

    # ---- 5. 原始数据记录表表头（名称/姓名现在就填，日期等实验后）----
    fill_record_head(doc, prof, spec['experiment'], spec.get('date') or '')

    doc.save(str(out_path))
    return out_path


def fill_record_head(doc, prof, experiment, date=''):
    """原始数据记录表表头：实验名称 / 实验日期 / 姓名"""
    t = find_table_by_cell(doc, '实验名称')
    vals = {(0, 1): experiment, (1, 1): date, (1, 3): prof['姓名']}
    for (r, c), v in vals.items():
        if not v:
            continue
        cell = t.cell(r, c)
        p = cell.paragraphs[0] if cell.paragraphs else cell.add_paragraph()
        for run in list(p.runs):                 # 别重复追加
            run._element.getparent().remove(run._element)
        style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_run(p.add_run(), v)
    return t


def _insert_figure(cell, img, caption, width_cm):
    """插图段落（居中、与下段同页）+ 题注段落（五号居中）"""
    # 图片段落
    p = cell.add_paragraph()
    style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.keep_with_next = True      # 图与题注不分离
    run = p.add_run()
    run.add_picture(str(img), width=Cm(width_cm))

    # 题注
    if caption:
        cp = cell.add_paragraph()
        style_para(cp, align=WD_ALIGN_PARAGRAPH.CENTER)
        style_run(cp.add_run(), caption, size=SZ_CAP)


# ==========================================================================
# 填原始数据（数据模式）
# ==========================================================================
def fill_data(spec, docx_path, profile=None):
    _need_docx()
    """
    spec = {
      "date": "9.18",
      "title": "最小偏向角测量数据",          # 数据表上方的小标题，可省略
      "header": [ [...每列文本或 [文本,是否下标]...] ],
      "rows":   [ [...], ... ],
      "widths": [800, 1440, ...]              # dxa，省略则均分
    }
    """
    prof = profile or read_profile()
    doc = docx.Document(str(docx_path))

    # ---- 1. 正文信息行的实验日期 ----
    if spec.get('date'):
        info1 = find_para(doc, '实验日期：')
        set_blank(underlined_runs(info1), [0], ' %s ' % spec['date'])

    # ---- 2. 原始数据记录表 ----
    t = fill_record_head(doc, prof, spec['experiment'], spec.get('date') or '')

    # 大格原为“至少 21cm”，加上表头后一页放不下，会把表头和数据拆到两页。
    # 数据是打进去的，按内容自适应即可，只留一个够体面的最小高度。
    row = t.rows[-1]
    trPr = row._tr.find(qn('w:trPr'))
    if trPr is not None:
        h = trPr.find(qn('w:trHeight'))
        if h is not None:
            h.set(qn('w:val'), '8500')      # 15 cm
            h.set(qn('w:hRule'), 'atLeast')

    big = t.cell(2, 0)
    clear_cell(big)
    if spec.get('title'):
        cell_para(big, spec['title'], bold=True,
                  align=WD_ALIGN_PARAGRAPH.LEFT, line=1.2)

    ncol = len(spec['header'])
    widths = spec.get('widths') or [int(8000 / ncol)] * ncol
    sub = big.add_table(rows=1 + len(spec['rows']), cols=ncol)
    sub.alignment = WD_TABLE_ALIGNMENT.CENTER
    sub.autofit = False
    set_table_borders(sub)

    def put(cell, parts):
        p = cell.paragraphs[0]
        style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
        if isinstance(parts, str):
            parts = [(parts, False)]
        for txt, is_sub in parts:
            style_run(p.add_run(), txt, size=SZ_TABLE, sub=is_sub)

    for j, col in enumerate(spec['header']):
        c = sub.cell(0, j)
        p = c.paragraphs[0]
        style_para(p, align=WD_ALIGN_PARAGRAPH.CENTER, line=1.0)
        parts = col if isinstance(col, list) else [(col, False)]
        for txt, is_sub in parts:
            style_run(p.add_run(), txt, size=SZ_TABLE, bold=True, sub=is_sub)
    for i, row in enumerate(spec['rows'], start=1):
        for j, v in enumerate(row):
            put(sub.cell(i, j), v)

    for j, w in enumerate(widths):
        for r in sub.rows:
            r.cells[j].width = Pt(w / 20)

    big.add_paragraph()

    # ---- 3. 原始数据记录表单独起页（21cm 大格会把表头挤到上一页）----
    try:
        find_para(doc, '原始数据记录表').paragraph_format.page_break_before = True
    except LookupError:
        pass

    doc.save(str(docx_path))
    return docx_path


# ==========================================================================
# 导出 PDF
# ==========================================================================
def export_pdf(docx_path, pdf_path=None, timeout=300):
    docx_path = Path(docx_path)
    pdf_path = Path(pdf_path) if pdf_path else docx_path.with_suffix('.pdf')
    outdir = pdf_path.parent
    outdir.mkdir(parents=True, exist_ok=True)
    if shutil.which('soffice') is None:
        raise RuntimeError(
            '找不到 soffice（LibreOffice）。导出 PDF 需要它：\n'
            '  Debian/Ubuntu:  sudo apt install libreoffice-writer\n'
            '  macOS:          brew install --cask libreoffice')
    try:
        subprocess.run(
            ['soffice', '--headless', '--convert-to', 'pdf',
             '--outdir', str(outdir), str(docx_path)],
            check=True, timeout=timeout,
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except subprocess.CalledProcessError as e:
        raise RuntimeError('LibreOffice 转换失败。若另一个 soffice 正在运行会互相锁死，'
                           '先 `pgrep -a soffice` 确认后重试。') from e
    produced = outdir / (docx_path.stem + '.pdf')
    if produced != pdf_path and produced.exists():
        shutil.move(str(produced), str(pdf_path))
    return pdf_path


def pdf_pages(pdf_path):
    r = subprocess.run(['pdfinfo', str(pdf_path)],
                       capture_output=True, text=True)
    m = re.search(r'Pages:\s+(\d+)', r.stdout)
    return int(m.group(1)) if m else 0


# ==========================================================================
# 自检
# ==========================================================================
def verify(docx_path, expect_figs=0, prof=None, experiment=None):
    _need_docx()
    """返回 (问题列表, 信息字典)。只检查正文表，封面用三号是正常的。"""
    problems, info = [], {}
    doc = docx.Document(str(docx_path))
    xml = doc.element.xml

    # 红色批注残留（模板理论上已清干净，防手滑）
    if 'FF0000' in xml.upper():
        problems.append('文档里出现红色(FF0000)文字')

    # 只盯正文表
    try:
        body_tbl = find_table_by_cell(doc, '实验目的')
    except LookupError:
        return ['找不到正文表（实验目的/仪器/原理）'], info

    cell = body_tbl.cell(0, 0)
    nfig = len(cell._tc.findall('.//' + qn('w:drawing')))
    info['插图数'] = nfig
    if expect_figs and nfig < expect_figs:
        problems.append('插图数 %d < 预期 %d' % (nfig, expect_figs))

    # 正文格式：跳过图片段和题注段
    CAP = tuple('图%d' % i for i in range(1, 10))
    bad = set()
    for p in cell.paragraphs:
        if p._element.findall('.//' + qn('w:drawing')):
            continue
        if p.text.strip().startswith(CAP):
            continue
        if p.paragraph_format.line_spacing != LINE:
            bad.add('行距')
        for r in p.runs:
            if not r.text.strip():
                continue
            if r.font.size and r.font.size.pt * 2 != SZ_BODY:
                bad.add('字号')
            if r.font.name != FONT:
                bad.add('字体')
    if bad:
        problems.append('正文格式不合规：%s' % '、'.join(sorted(bad)))

    # 必填信息
    if prof:
        txt = '\n'.join(p.text for p in doc.paragraphs)
        for k in ('姓名', '座位号'):
            if prof.get(k) and prof[k] not in txt:
                problems.append('信息行缺少 %s' % k)
    if experiment:
        head = '\n'.join(p.text for p in doc.paragraphs)
        if experiment not in head:
            problems.append('标题里没有实验名称')

    return problems, info


# ==========================================================================
# 环境自检（lab.py doctor）
# ==========================================================================
def check_environment():
    """返回 (致命问题, 提醒, 信息行)。致命问题会导致 build/export 失败。"""
    fatal, warn, info = [], [], []

    # 依赖。注意模块名 ≠ pip 包名（docx 要装 python-docx，装成 docx 是另一个废弃库）
    DEPS = [('docx',       'python-docx', '读写 docx',      True),
            ('lxml',       'lxml',        'python-docx 依赖', True),
            ('PIL',        'Pillow',      '读图片尺寸',      False),
            ('matplotlib', 'matplotlib',  '画原理图',        False)]
    for mod, pkg, why, hard in DEPS:
        try:
            __import__(mod)
        except BaseException:      # 二进制依赖损坏时可能抛 ImportError 以外的东西
            (fatal if hard else warn).append(
                '缺少或损坏 %s（%s）：pip install %s' % (mod, why, pkg))

    # 目录与模板
    info.append('工作根目录 PHYSLAB_HOME : %s' % ROOT)
    info.append('模板       PHYSLAB_TEMPLATE: %s' % TEMPLATE)
    if not TEMPLATE.exists():
        fatal.append(
            '找不到空白模板：%s\n'
            '      本仓库不附带任何学校的模板文件（版权原因），需要你自己准备：\n'
            '      1) 拿到本校的实验报告模板\n'
            '      2) 删掉里面的示例/批注文字，只留空白表单\n'
            '      3) 放到上面这个路径，或设 PHYSLAB_TEMPLATE 指向它\n'
            '      详见 reference/format.md' % TEMPLATE)
    if not PROFILE.exists():
        warn.append('还没有个人信息文件：%s（可用 lab.py profile 查看）' % PROFILE)
    if not os.access(ROOT if ROOT.exists() else ROOT.parent, os.W_OK):
        fatal.append('目录不可写：%s' % ROOT)

    # 外部工具
    info.append('LibreOffice(soffice)     : %s' % (shutil.which('soffice') or '未找到'))
    if shutil.which('soffice') is None:
        fatal.append('缺少 soffice，无法导出 PDF：apt install libreoffice-writer')

    return fatal, warn, info
